"""PPT 配音模块单元测试。

覆盖主要功能点、边界条件和异常情况，
使用 Python 标准 unittest 框架。
"""

import asyncio
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, PropertyMock, patch

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from edge_tts.exceptions import (
    NoAudioReceived,
    UnexpectedResponse,
    UnknownResponse,
    WebSocketError,
)
from lxml import etree as lxml_etree
from pptx import Presentation
from pptx.util import Inches

from ppt_speech import (
    PptSpeechConfig,
    clear_advance_after_time,
    embed_audio_autoplay,
    get_audio_duration,
    normalize_voice_name,
    process_slides,
    read_notes_text,
    remove_embedded_audio,
    set_advance_after_time,
    speak_ppt_notes,
    text_to_mp3,
)
from ppt_speech.audio import P14_NS, P_NS, _apply_autoplay_timing
from ppt_speech.slide_transition import _clear_adv_tm, _set_adv_tm


NSMAP = {"p": P_NS, "p14": P14_NS}


class TestPTSpeechConfig(unittest.TestCase):
    """PTSpeechConfig 配置类测试。"""

    def test_default_creation(self) -> None:
        """测试默认配置创建。"""
        config = PptSpeechConfig()
        self.assertEqual(config.input_dir, Path("data"))
        self.assertEqual(config.output_dir, Path("data"))
        self.assertEqual(config.input_filename, "input.pptx")
        self.assertEqual(config.output_filename, "output.pptx")
        self.assertEqual(config.voice_name, "zh-CN-XiaoxiaoNeural")
        self.assertEqual(config.speech_rate, "+0%")
        self.assertEqual(config.audio_icon_offset, -2.0)
        self.assertEqual(config.audio_icon_size, 1.0)
        self.assertTrue(config.auto_advance)
        self.assertEqual(config.auto_advance_delay, 2.0)

    def test_custom_creation(self) -> None:
        """测试自定义配置创建。"""
        config = PptSpeechConfig(
            input_dir=Path("input_dir"),
            output_dir=Path("output_dir"),
            input_filename="test.pptx",
            output_filename="result.pptx",
            voice_name="en-US-AriaNeural",
            speech_rate="-50%",
            temp_audio_dir=Path("custom_temp"),
            audio_icon_offset=-3.0,
            audio_icon_size=2.0,
        )
        self.assertEqual(config.input_dir, Path("input_dir"))
        self.assertEqual(config.voice_name, "en-US-AriaNeural")
        self.assertEqual(config.speech_rate, "-50%")
        self.assertEqual(config.audio_icon_offset, -3.0)
        self.assertEqual(config.audio_icon_size, 2.0)

    def test_input_path_property(self) -> None:
        """测试 input_path 属性。"""
        config = PptSpeechConfig(
            input_dir=Path("mydata"),
            input_filename="test.pptx",
        )
        self.assertEqual(config.input_path, Path("mydata/test.pptx"))

    def test_output_path_property(self) -> None:
        """测试 output_path 属性。"""
        config = PptSpeechConfig(
            output_dir=Path("out"),
            output_filename="result.pptx",
        )
        self.assertEqual(config.output_path, Path("out/result.pptx"))

    @patch("pathlib.Path.exists", return_value=True)
    def test_validate_valid(self, mock_exists: MagicMock) -> None:
        """测试合法配置验证。"""
        config = PptSpeechConfig()
        config.validate()

    def test_validate_invalid_voice(self) -> None:
        """测试无效语音名称验证。"""
        config = PptSpeechConfig(voice_name="invalid-voice")
        with self.assertRaises(ValueError) as ctx:
            config.validate()
        self.assertIn("语音名称格式错误", str(ctx.exception))

    @patch("pathlib.Path.exists", return_value=True)
    def test_validate_invalid_speech_rate(self, mock_exists: MagicMock) -> None:
        """测试无效语速验证。"""
        config = PptSpeechConfig(speech_rate="fast")
        with self.assertRaises(ValueError) as ctx:
            config.validate()
        self.assertIn("语速格式错误", str(ctx.exception))

    def test_validate_missing_input_file(self) -> None:
        """测试输入文件不存在验证。"""
        config = PptSpeechConfig(input_dir=Path("nonexistent"))
        with self.assertRaises(FileNotFoundError) as ctx:
            config.validate()
        self.assertIn("输入 PPT 文件不存在", str(ctx.exception))

    @patch("pathlib.Path.exists", return_value=True)
    def test_validate_boundary_speech_rates(self, mock_exists: MagicMock) -> None:
        """测试边界语速值验证。"""
        for rate in ("+0%", "-0%", "+100%", "-100%", "+999%"):
            config = PptSpeechConfig(speech_rate=rate)
            config.validate()

    @patch("pathlib.Path.exists", return_value=True)
    def test_validate_negative_auto_advance_delay(
        self, mock_exists: MagicMock
    ) -> None:
        """测试自动翻页延迟为负数时验证失败。"""
        config = PptSpeechConfig(auto_advance=True, auto_advance_delay=-1.0)
        with self.assertRaises(ValueError) as ctx:
            config.validate()
        self.assertIn("自动翻页延迟时间不能为负数", str(ctx.exception))


class TestNormalizeVoiceName(unittest.TestCase):
    """_normalize_voice_name 函数测试。"""

    def test_standard_chinese_voice(self) -> None:
        """测试标准中文语音名称规范化。"""
        result = normalize_voice_name("zh-CN-XiaoxiaoNeural")
        self.assertEqual(
            result,
            "Microsoft Server Speech Text to Speech Voice"
            " (zh-CN, XiaoxiaoNeural)",
        )

    def test_standard_english_voice(self) -> None:
        """测试标准英文语音名称规范化。"""
        result = normalize_voice_name("en-US-AriaNeural")
        self.assertEqual(
            result,
            "Microsoft Server Speech Text to Speech Voice"
            " (en-US, AriaNeural)",
        )

    def test_voice_with_region_suffix(self) -> None:
        """测试带地区后缀的语音名称规范化。"""
        result = normalize_voice_name("zh-CN-Xiaoxiao-YunxiNeural")
        self.assertEqual(
            result,
            "Microsoft Server Speech Text to Speech Voice"
            " (zh-CN-Xiaoxiao, YunxiNeural)",
        )

    def test_already_full_format(self) -> None:
        """测试已为完整格式的语音名称。"""
        full_name = (
            "Microsoft Server Speech Text to Speech Voice"
            " (zh-CN, XiaoxiaoNeural)"
        )
        result = normalize_voice_name(full_name)
        self.assertEqual(result, full_name)

    def test_non_standard_voice(self) -> None:
        """测试非标准格式的语音名称。"""
        result = normalize_voice_name("SomeVoice")
        self.assertEqual(result, "SomeVoice")


class TestReadNotesText(unittest.TestCase):
    """_read_notes_text 函数测试。"""

    def setUp(self) -> None:
        self.mock_slide = MagicMock()

    def test_slide_with_notes(self) -> None:
        """测试有备注的幻灯片。"""
        mock_notes = MagicMock()
        mock_text_frame = MagicMock()
        mock_text_frame.text = "这是一段备注文字"
        mock_notes.notes_text_frame = mock_text_frame

        self.mock_slide.has_notes_slide = True
        self.mock_slide.notes_slide = mock_notes

        result = read_notes_text(self.mock_slide)
        self.assertEqual(result, "这是一段备注文字")

    def test_slide_without_notes_slide(self) -> None:
        """测试无备注页的幻灯片。"""
        self.mock_slide.has_notes_slide = False
        result = read_notes_text(self.mock_slide)
        self.assertEqual(result, "")

    def test_slide_with_empty_text_frame(self) -> None:
        """测试备注文本框为空的幻灯片。"""
        mock_notes = MagicMock()
        mock_notes.notes_text_frame = None

        self.mock_slide.has_notes_slide = True
        self.mock_slide.notes_slide = mock_notes

        result = read_notes_text(self.mock_slide)
        self.assertEqual(result, "")

    def test_slide_with_whitespace_only_notes(self) -> None:
        """测试备注仅含空白字符的幻灯片。"""
        mock_notes = MagicMock()
        mock_text_frame = MagicMock()
        mock_text_frame.text = "   \n\t   "
        mock_notes.notes_text_frame = mock_text_frame

        self.mock_slide.has_notes_slide = True
        self.mock_slide.notes_slide = mock_notes

        result = read_notes_text(self.mock_slide)
        self.assertEqual(result, "")


class TestTextToMP3(unittest.IsolatedAsyncioTestCase):
    """text_to_mp3 函数测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    async def test_empty_text_returns_false(self) -> None:
        """测试空文字返回 False。"""
        result = await text_to_mp3("", self.temp_dir / "test.mp3")
        self.assertFalse(result)

    async def test_whitespace_only_returns_false(self) -> None:
        """测试纯空白文字返回 False。"""
        result = await text_to_mp3(
            "   \n\t   ", self.temp_dir / "test.mp3"
        )
        self.assertFalse(result)

    async def test_invalid_speech_rate_raises(self) -> None:
        """测试无效语速抛出 ValueError。"""
        with self.assertRaises(ValueError) as ctx:
            await text_to_mp3(
                "测试文字",
                self.temp_dir / "test.mp3",
                speech_rate="invalid",
            )
        self.assertIn("语速格式错误", str(ctx.exception))

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_successful_conversion(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试成功转换。"""
        mock_comm = AsyncMock()
        mock_communicate_class.return_value = mock_comm

        result = await text_to_mp3(
            "测试文字",
            self.temp_dir / "test.mp3",
            voice_name="zh-CN-XiaoxiaoNeural",
            speech_rate="+50%",
        )

        self.assertTrue(result)
        mock_communicate_class.assert_called_once()
        mock_comm.save.assert_called_once()

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_no_audio_received_handling(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试 NoAudioReceived 异常处理。"""
        from edge_tts.exceptions import EdgeTTSException

        mock_comm = AsyncMock()
        mock_comm.save.side_effect = NoAudioReceived("无音频")
        mock_communicate_class.return_value = mock_comm

        with self.assertRaises(EdgeTTSException):
            await text_to_mp3("测试文字", self.temp_dir / "test.mp3")

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_unexpected_response_handling(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试 UnexpectedResponse 异常处理。"""
        from edge_tts.exceptions import EdgeTTSException

        mock_comm = AsyncMock()
        mock_comm.save.side_effect = UnexpectedResponse("意外响应")
        mock_communicate_class.return_value = mock_comm

        with self.assertRaises(EdgeTTSException):
            await text_to_mp3("测试文字", self.temp_dir / "test.mp3")

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_unknown_response_handling(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试 UnknownResponse 异常处理。"""
        from edge_tts.exceptions import EdgeTTSException

        mock_comm = AsyncMock()
        mock_comm.save.side_effect = UnknownResponse("未知响应")
        mock_communicate_class.return_value = mock_comm

        with self.assertRaises(EdgeTTSException):
            await text_to_mp3("测试文字", self.temp_dir / "test.mp3")

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_websocket_error_handling(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试 WebSocketError 异常处理。"""
        from edge_tts.exceptions import EdgeTTSException

        mock_comm = AsyncMock()
        mock_comm.save.side_effect = WebSocketError("连接错误")
        mock_communicate_class.return_value = mock_comm

        with self.assertRaises(EdgeTTSException):
            await text_to_mp3("测试文字", self.temp_dir / "test.mp3")

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_os_error_handling(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试 OS 错误处理。"""
        mock_comm = AsyncMock()
        mock_comm.save.side_effect = OSError("磁盘错误")
        mock_communicate_class.return_value = mock_comm

        with self.assertRaises(OSError) as ctx:
            await text_to_mp3("测试文字", self.temp_dir / "test.mp3")
        self.assertIn("磁盘错误", str(ctx.exception))

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_custom_voice_and_rate(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试自定义语音和语速。"""
        mock_comm = AsyncMock()
        mock_communicate_class.return_value = mock_comm

        result = await text_to_mp3(
            "测试",
            self.temp_dir / "test.mp3",
            voice_name="en-US-AriaNeural",
            speech_rate="-30%",
        )

        self.assertTrue(result)
        call_args = mock_communicate_class.call_args
        self.assertIn("AriaNeural", call_args[0][1])
        self.assertEqual(call_args[1]["rate"], "-30%")

    @patch("ppt_speech.core.tts_client.Communicate")
    async def test_creates_parent_directory(
        self, mock_communicate_class: MagicMock
    ) -> None:
        """测试自动创建父目录。"""
        mock_comm = AsyncMock()
        mock_communicate_class.return_value = mock_comm

        deep_path = self.temp_dir / "a" / "b" / "c" / "test.mp3"
        await text_to_mp3("测试", deep_path)

        self.assertTrue(deep_path.parent.exists())


class TestApplyAutoplayTiming(unittest.TestCase):
    """_apply_autoplay_timing 函数测试。"""

    def test_xml_with_timing(self) -> None:
        """测试有时序 XML 的修改。"""
        xml_str = (
            f'<p:sld xmlns:p="{P_NS}" xmlns:p14="{P14_NS}">'
            f"<p:timing>"
            f"<p:tnLst>"
            f"<p:par>"
            f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
            f"<p:childTnLst>"
            f'<p:seq concurrent="1" nextAc="seek">'
            f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
            f"<p:childTnLst>"
            f"<p:par>"
            f'<p:cMediaNode volCtrl="100" muteVal="0">'
            f'<p:cTn id="3" fill="hold">'
            f"<p:stCondLst>"
            f'<p:cond delay="0"/>'
            f"</p:stCondLst>"
            f"</p:cTn>"
            f"</p:cMediaNode>"
            f"</p:par>"
            f"</p:childTnLst>"
            f"</p:cTn>"
            f"</p:seq>"
            f"</p:childTnLst>"
            f"</p:cTn>"
            f"</p:par>"
            f"</p:tnLst>"
            f"</p:timing>"
            f"</p:sld>"
        )

        element = lxml_etree.fromstring(xml_str.encode())
        _apply_autoplay_timing(element)

        media_nodes = element.findall(".//p:cMediaNode", NSMAP)
        self.assertTrue(len(media_nodes) > 0)
        for node in media_nodes:
            self.assertEqual(
                node.get(f"{{{P_NS}}}playOnEntry"), "1"
            )

        conds = element.findall(".//p:cond", NSMAP)
        for cond in conds:
            self.assertEqual(cond.get("delay"), "0")
            with_prev = cond.find(f"{{{P_NS}}}withPrev")
            self.assertIsNotNone(with_prev)

    def test_xml_without_timing(self) -> None:
        """测试无时序 XML 的处理。"""
        xml_str = f'<p:sld xmlns:p="{P_NS}"><p:someChild/></p:sld>'
        element = lxml_etree.fromstring(xml_str.encode())

        _apply_autoplay_timing(element)

        self.assertEqual(element.tag.split("}")[-1], "sld")

    def test_xml_with_media_elements(self) -> None:
        """测试媒体元素的 playOnEntry 设置。"""
        xml_str = (
            f'<p:sld xmlns:p="{P_NS}" xmlns:p14="{P14_NS}">'
            f"<p:p14BlipFill>"
            f"<p14:media>"
            f"</p14:media>"
            f"</p:p14BlipFill>"
            f"</p:sld>"
        )

        element = lxml_etree.fromstring(xml_str.encode())
        _apply_autoplay_timing(element)

        media_elements = list(element.iter(f"{{{P14_NS}}}media"))
        self.assertTrue(len(media_elements) > 0)
        for media_el in media_elements:
            self.assertEqual(
                media_el.get(f"{{{P_NS}}}playOnEntry"), "1"
            )


class TestEmbedAudioAutoplay(unittest.TestCase):
    """embed_audio_autoplay 函数测试。"""

    def setUp(self) -> None:
        self.mock_slide = MagicMock()
        self.mock_slide._element = MagicMock()
        self.mock_slide.shapes = MagicMock()

        self.temp_dir = Path(tempfile.mkdtemp())
        self.audio_path = self.temp_dir / "test.mp3"
        self.audio_path.write_bytes(b"fake audio content")

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_successful_embedding(self) -> None:
        """测试成功嵌入音频。"""
        embed_audio_autoplay(self.mock_slide, self.audio_path)

        self.mock_slide.shapes.add_movie.assert_called_once()
        call_args = self.mock_slide.shapes.add_movie.call_args
        self.assertIn(str(self.audio_path), call_args[0])

    def test_file_not_found(self) -> None:
        """测试音频文件不存在异常。"""
        nonexistent_path = self.temp_dir / "nonexistent.mp3"
        with self.assertRaises(FileNotFoundError) as ctx:
            embed_audio_autoplay(self.mock_slide, nonexistent_path)
        self.assertIn("音频文件不存在", str(ctx.exception))

    def test_invalid_icon_size(self) -> None:
        """测试无效图标尺寸异常。"""
        with self.assertRaises(ValueError) as ctx:
            embed_audio_autoplay(
                self.mock_slide, self.audio_path, icon_size=0
            )
        self.assertIn("图标尺寸必须为正数", str(ctx.exception))

    def test_negative_icon_size_raises(self) -> None:
        """测试负图标尺寸异常。"""
        with self.assertRaises(ValueError):
            embed_audio_autoplay(
                self.mock_slide, self.audio_path, icon_size=-1.0
            )

    def test_custom_icon_parameters(self) -> None:
        """测试自定义图标参数。"""
        embed_audio_autoplay(
            self.mock_slide,
            self.audio_path,
            icon_offset=-3.0,
            icon_size=2.0,
        )

        call_kwargs = self.mock_slide.shapes.add_movie.call_args[1]
        self.assertIn("left", call_kwargs)
        self.assertIn("width", call_kwargs)


class TestRemoveEmbeddedAudio(unittest.TestCase):
    """remove_embedded_audio 函数测试（真实 python-pptx 幻灯片）。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.audio_path = self.temp_dir / "slide_1.mp3"
        self.audio_path.write_bytes(b"ID3" + b"\x00" * 200)
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _pic_count(self) -> int:
        return len(self.slide._element.findall(f".//{{{P_NS}}}pic"))

    def _video_timing_count(self) -> int:
        return len(self.slide._element.findall(f".//{{{P_NS}}}video"))

    def test_plain_slide_removes_nothing(self) -> None:
        """空白幻灯片清理返回 0，且不产生副作用。"""
        timing_before = self.slide._element.find(f"{{{P_NS}}}timing")
        self.assertEqual(remove_embedded_audio(self.slide), 0)
        self.assertEqual(self._pic_count(), 0)
        self.assertEqual(
            self.slide._element.find(f"{{{P_NS}}}timing") is not None,
            timing_before is not None,
        )

    def test_removes_shape_timing_and_rels(self) -> None:
        """嵌入后清理：形状、时序、媒体/图片关系全部移除。"""
        embed_audio_autoplay(self.slide, self.audio_path)
        self.assertEqual(self._pic_count(), 1)

        removed = remove_embedded_audio(self.slide)

        self.assertEqual(removed, 1)
        self.assertEqual(self._pic_count(), 0)
        self.assertEqual(self._video_timing_count(), 0)
        self.assertIsNone(
            self.slide._element.find(f"{{{P_NS}}}timing"),
            "空骨架 p:timing 应整体移除",
        )
        rel_suffixes = {
            rel.reltype.rsplit("/", 1)[-1]
            for rel in self.slide.part.rels.values()
            if not rel.is_external
        }
        self.assertNotIn("media", rel_suffixes)
        self.assertNotIn("video", rel_suffixes)
        self.assertNotIn("image", rel_suffixes)

    def test_media_parts_dropped_from_saved_package(self) -> None:
        """清理并保存后，包内不再保留旧 MP3/海报帧部件。"""
        embed_audio_autoplay(self.slide, self.audio_path)
        remove_embedded_audio(self.slide)

        out_path = self.temp_dir / "stripped.pptx"
        self.prs.save(str(out_path))

        with zipfile.ZipFile(out_path) as archive:
            media_parts = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/media/")
            ]
        self.assertEqual(media_parts, [])

    def test_legacy_unmarked_audio_is_recognized(self) -> None:
        """无标记的历史版本产物：经保存重开后仍能按音频部件识别并清理。"""
        embed_audio_autoplay(self.slide, self.audio_path)
        cnvpr = self.slide._element.find(
            f".//{{{P_NS}}}pic/{{{P_NS}}}nvPicPr/{{{P_NS}}}cNvPr"
        )
        del cnvpr.attrib["descr"]
        legacy_path = self.temp_dir / "legacy.pptx"
        self.prs.save(str(legacy_path))

        reopened = Presentation(str(legacy_path))
        reopened_slide = reopened.slides[0]
        self.assertEqual(remove_embedded_audio(reopened_slide), 1)
        self.assertEqual(
            len(reopened_slide._element.findall(f".//{{{P_NS}}}pic")), 0
        )

        redubbed_path = self.temp_dir / "redubbed.pptx"
        reopened.save(str(redubbed_path))
        with zipfile.ZipFile(redubbed_path) as archive:
            media_parts = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/media/")
            ]
        self.assertEqual(media_parts, [])

    def test_user_video_is_not_removed(self) -> None:
        """用户自行插入的 video/mp4 不应被当作旧配音删除。"""
        video_path = self.temp_dir / "user.mp4"
        video_path.write_bytes(b"\x00" * 200)
        self.slide.shapes.add_movie(
            str(video_path),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            poster_frame_image=None,
            mime_type="video/mp4",
        )

        self.assertEqual(remove_embedded_audio(self.slide), 0)
        self.assertEqual(self._pic_count(), 1)
        self.assertEqual(self._video_timing_count(), 1)
        self.assertIsNotNone(
            self.slide._element.find(f"{{{P_NS}}}timing")
        )

    def test_keeps_user_video_when_stripping_dubbing(self) -> None:
        """同页用户视频与本工具配音共存时，只删除配音，保留视频与时序。"""
        video_path = self.temp_dir / "user.mp4"
        video_path.write_bytes(b"\x00" * 200)
        self.slide.shapes.add_movie(
            str(video_path),
            left=Inches(1),
            top=Inches(1),
            width=Inches(2),
            height=Inches(2),
            poster_frame_image=None,
            mime_type="video/mp4",
        )
        embed_audio_autoplay(self.slide, self.audio_path)
        self.assertEqual(self._pic_count(), 2)

        removed = remove_embedded_audio(self.slide)

        self.assertEqual(removed, 1)
        self.assertEqual(self._pic_count(), 1)
        self.assertEqual(self._video_timing_count(), 1)
        self.assertIsNotNone(
            self.slide._element.find(f"{{{P_NS}}}timing")
        )


class TestEmbedAudioAutoplayIdempotency(unittest.TestCase):
    """embed_audio_autoplay 重复调用的幂等性测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.audio_path = self.temp_dir / "slide_1.mp3"
        self.audio_path.write_bytes(b"ID3" + b"\x00" * 200)
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_repeated_embed_does_not_accumulate(self) -> None:
        """连续嵌入三次：始终只有一个形状、一个时序节点，且带标记。"""
        removed_first = embed_audio_autoplay(self.slide, self.audio_path)
        removed_second = embed_audio_autoplay(self.slide, self.audio_path)
        removed_third = embed_audio_autoplay(self.slide, self.audio_path)

        self.assertEqual(removed_first, 0)
        self.assertEqual(removed_second, 1)
        self.assertEqual(removed_third, 1)

        pics = self.slide._element.findall(f".//{{{P_NS}}}pic")
        videos = self.slide._element.findall(f".//{{{P_NS}}}video")
        self.assertEqual(len(pics), 1)
        self.assertEqual(len(videos), 1)

        cnvpr = pics[0].find(f"{{{P_NS}}}nvPicPr/{{{P_NS}}}cNvPr")
        self.assertEqual(cnvpr.get("descr"), "ppt-speech:dubbing-audio")

    def test_repeated_embed_keeps_single_media_part(self) -> None:
        """重复嵌入保存后包内仅保留一份 MP3 媒体部件。"""
        embed_audio_autoplay(self.slide, self.audio_path)
        embed_audio_autoplay(self.slide, self.audio_path)
        embed_audio_autoplay(self.slide, self.audio_path)

        out_path = self.temp_dir / "idempotent.pptx"
        self.prs.save(str(out_path))

        with zipfile.ZipFile(out_path) as archive:
            mp3_parts = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/media/") and name.endswith(".mp3")
            ]
        self.assertEqual(len(mp3_parts), 1)

    def test_validation_runs_before_cleanup(self) -> None:
        """参数校验失败时不应先破坏既有配音。"""
        embed_audio_autoplay(self.slide, self.audio_path)
        with self.assertRaises(ValueError):
            embed_audio_autoplay(self.slide, self.audio_path, icon_size=0)
        self.assertEqual(
            len(self.slide._element.findall(f".//{{{P_NS}}}pic")), 1
        )

    def test_redub_saved_file_in_new_session(self) -> None:
        """已配音文件保存后重新打开并再次嵌入（回归：不得抛 AttributeError）。

        python-pptx 1.0.x 的 PartFactory 未注册 audio/mpeg，重载后音频
        部件成为无 sha1 的基类 Part；add_movie 会扫描包内**所有**媒体
        关系，因此重配任意一页时，其它页面残留的音频部件即触发
        ``AttributeError: 'Part' object has no attribute 'sha1'``。
        需要至少两页才能复现（重配页自身的旧关系已被先清理）。
        """
        # 使用独立演示文稿，避免混入 setUp 中已有的空白页
        prs = Presentation()
        slides = [
            prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(2)
        ]
        audio_paths = []
        for idx, slide in enumerate(slides):
            audio_path = self.temp_dir / f"slide_{idx}.mp3"
            audio_path.write_bytes(b"ID3" + bytes([idx]) * 200)
            audio_paths.append(audio_path)
            embed_audio_autoplay(slide, audio_path)

        first_path = self.temp_dir / "first.pptx"
        prs.save(str(first_path))

        # 模拟下一次配音：重新打开文件，仅重配第 1 页（第 2 页关系仍在）
        reopened = Presentation(str(first_path))
        reopened_slide = reopened.slides[0]
        # 修复前此行即抛出 AttributeError
        removed = embed_audio_autoplay(reopened_slide, audio_paths[0])
        self.assertEqual(removed, 1)

        second_path = self.temp_dir / "second.pptx"
        reopened.save(str(second_path))

        with zipfile.ZipFile(second_path) as archive:
            mp3_parts = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/media/") and name.endswith(".mp3")
            ]
        self.assertEqual(len(mp3_parts), 2)

    def test_audio_content_types_registered_as_media_part(self) -> None:
        """常见音频 content-type 已注册到 PartFactory，重载即为 MediaPart。"""
        from pptx.opc.package import PartFactory
        from pptx.parts.media import MediaPart

        self.assertIs(
            PartFactory.part_type_for["audio/mpeg"], MediaPart
        )


class TestProcessSlides(unittest.IsolatedAsyncioTestCase):
    """process_slides 函数测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.config = PptSpeechConfig(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
            temp_audio_dir=self.temp_dir / "temp_audio",
            auto_advance=False,
        )

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_process_with_notes(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试处理有备注的幻灯片。"""
        mock_read_notes.return_value = "测试备注文字"
        mock_text_to_mp3.return_value = True

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, self.config)

        mock_text_to_mp3.assert_called_once_with(
            "测试备注文字",
            self.config.temp_audio_dir / "slide_1.mp3",
            voice_name=self.config.voice_name,
            speech_rate=self.config.speech_rate,
        )
        mock_embed.assert_called_once()
        mock_prs.save.assert_called_once()

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_process_without_notes(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试处理无备注的幻灯片。"""
        mock_read_notes.return_value = ""

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, self.config)

        mock_text_to_mp3.assert_not_called()
        mock_embed.assert_not_called()
        mock_prs.save.assert_called_once()

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_process_multiple_slides(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试处理多个幻灯片。"""
        mock_read_notes.side_effect = [
            "第一页备注",
            "",
            "第三页备注",
        ]
        mock_text_to_mp3.return_value = True

        mock_slides = [MagicMock(), MagicMock(), MagicMock()]
        mock_prs = MagicMock()
        mock_prs.slides = mock_slides

        await process_slides(mock_prs, self.config)

        self.assertEqual(mock_text_to_mp3.call_count, 2)
        self.assertEqual(mock_embed.call_count, 2)
        mock_prs.save.assert_called_once()

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_temp_dir_cleanup_on_success(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试成功处理后临时目录清理。"""
        mock_read_notes.return_value = "测试"
        mock_text_to_mp3.return_value = True

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        self.config.temp_audio_dir.mkdir(parents=True, exist_ok=True)

        await process_slides(mock_prs, self.config)

        self.assertFalse(self.config.temp_audio_dir.exists())

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_temp_dir_cleanup_on_failure(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试处理失败后临时目录清理。"""
        mock_read_notes.return_value = "测试"
        mock_text_to_mp3.side_effect = RuntimeError("处理失败")

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        self.config.temp_audio_dir.mkdir(parents=True, exist_ok=True)

        with self.assertRaises(RuntimeError):
            await process_slides(mock_prs, self.config)

        self.assertFalse(self.config.temp_audio_dir.exists())

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_tts_failure_skips_embed(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """测试 TTS 失败时跳过嵌入。"""
        mock_read_notes.return_value = "测试"
        mock_text_to_mp3.return_value = False

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, self.config)

        mock_embed.assert_not_called()

    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_default_temp_dir_uses_system_tempfile(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
    ) -> None:
        """temp_audio_dir 为 None 时应使用系统临时目录并自动清理。"""
        mock_read_notes.return_value = "测试备注"
        mock_text_to_mp3.return_value = True

        # 不指定 temp_audio_dir，验证默认走系统临时目录（tempfile）路径
        config = PptSpeechConfig(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
            auto_advance=False,
        )

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, config)

        # text_to_mp3 的第二个位置参数即为音频文件路径
        self.assertEqual(mock_text_to_mp3.call_count, 1)
        used_dir = mock_text_to_mp3.call_args.args[1].parent
        # 临时目录应位于系统临时目录下，而非当前工作目录的 .tmp_audio
        self.assertTrue(
            str(used_dir).startswith(str(Path(tempfile.gettempdir()))),
            f"临时目录 {used_dir} 不在系统临时目录下",
        )
        # 处理完成后该临时目录应已被上下文管理器自动清理
        self.assertFalse(used_dir.exists())


class TestGetAudioDuration(unittest.TestCase):
    """get_audio_duration 函数测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_file_not_found(self) -> None:
        """测试音频文件不存在异常。"""
        nonexistent = self.temp_dir / "no.mp3"
        with self.assertRaises(FileNotFoundError) as ctx:
            get_audio_duration(nonexistent)
        self.assertIn("音频文件不存在", str(ctx.exception))

    @patch("ppt_speech.core.audio.duration.TinyTag")
    def test_successful(self, mock_tinytag: MagicMock) -> None:
        """测试成功读取音频时长。"""
        mock_tag = MagicMock()
        mock_tag.duration = 3.5
        mock_tinytag.get.return_value = mock_tag

        audio = self.temp_dir / "a.mp3"
        audio.write_bytes(b"fake")

        self.assertEqual(get_audio_duration(audio), 3.5)
        mock_tinytag.get.assert_called_once_with(str(audio))

    @patch("ppt_speech.core.audio.duration.TinyTag")
    def test_duration_none_raises(self, mock_tinytag: MagicMock) -> None:
        """测试时长为 None 时抛出 ValueError。"""
        mock_tag = MagicMock()
        mock_tag.duration = None
        mock_tinytag.get.return_value = mock_tag

        audio = self.temp_dir / "a.mp3"
        audio.write_bytes(b"fake")

        with self.assertRaises(ValueError) as ctx:
            get_audio_duration(audio)
        self.assertIn("无法获取音频时长", str(ctx.exception))

    @patch("ppt_speech.core.audio.duration.TinyTag")
    def test_tinytag_exception_raises_value_error(
        self, mock_tinytag: MagicMock
    ) -> None:
        """测试 tinytag 抛出异常时转换为 ValueError。"""
        mock_tinytag.get.side_effect = Exception("解析错误")

        audio = self.temp_dir / "a.mp3"
        audio.write_bytes(b"fake")

        with self.assertRaises(ValueError) as ctx:
            get_audio_duration(audio)
        self.assertIn("无法解析音频文件", str(ctx.exception))


class TestSetAdvanceAfterTime(unittest.TestCase):
    """set_advance_after_time / _set_adv_tm 测试。"""

    def test_sets_advtm_and_inserts_before_timing(self) -> None:
        """测试写入 advTm 且位于 p:timing 之前（schema 顺序）。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/><p:timing/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())

        _set_adv_tm(el, 7000)

        trans = el.find(f"{{{P_NS}}}transition")
        self.assertIsNotNone(trans)
        self.assertEqual(trans.get("advTm"), "7000")
        children = list(el)
        self.assertEqual(children[1].tag, f"{{{P_NS}}}transition")
        self.assertEqual(children[2].tag, f"{{{P_NS}}}timing")

    def test_appends_when_no_timing(self) -> None:
        """测试无 timing/extLst 时追加到末尾。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())

        _set_adv_tm(el, 1000)

        trans = el.find(f"{{{P_NS}}}transition")
        self.assertIsNotNone(trans)
        self.assertEqual(trans.get("advTm"), "1000")

    def test_preserves_existing_transition_effects(self) -> None:
        """测试已有切换效果时仅更新 advTm，保留子元素。"""
        xml = (
            f'<p:sld xmlns:p="{P_NS}"><p:cSld/>'
            f'<p:transition><p:fade/></p:transition><p:timing/></p:sld>'
        )
        el = lxml_etree.fromstring(xml.encode())

        _set_adv_tm(el, 3000)

        transitions = el.findall(f"{{{P_NS}}}transition")
        self.assertEqual(len(transitions), 1)
        self.assertEqual(transitions[0].get("advTm"), "3000")
        self.assertIsNotNone(transitions[0].find(f"{{{P_NS}}}fade"))

    def test_negative_delay_raises(self) -> None:
        """测试负延迟抛出 ValueError。"""
        with self.assertRaises(ValueError):
            set_advance_after_time(MagicMock(), -1.0)


class TestClearAdvanceAfterTime(unittest.TestCase):
    """clear_advance_after_time / _clear_adv_tm 测试。"""

    def test_clear_removes_bare_transition(self) -> None:
        """仅承载 advTm 的空 transition 应随 advTm 一并移除。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/><p:timing/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())
        _set_adv_tm(el, 5000)
        self.assertIsNotNone(el.find(f"{{{P_NS}}}transition"))

        changed = _clear_adv_tm(el)

        self.assertTrue(changed)
        self.assertIsNone(el.find(f"{{{P_NS}}}transition"))
        # schema 顺序未被破坏
        self.assertEqual(
            [child.tag for child in el],
            [f"{{{P_NS}}}cSld", f"{{{P_NS}}}timing"],
        )

    def test_clear_keeps_transition_effect(self) -> None:
        """带切换效果的 transition 仅删除 advTm，保留效果子元素。"""
        xml = (
            f'<p:sld xmlns:p="{P_NS}"><p:cSld/>'
            f'<p:transition advTm="3000"><p:fade/></p:transition>'
            f"<p:timing/></p:sld>"
        )
        el = lxml_etree.fromstring(xml.encode())

        changed = _clear_adv_tm(el)

        self.assertTrue(changed)
        transitions = el.findall(f"{{{P_NS}}}transition")
        self.assertEqual(len(transitions), 1)
        self.assertIsNone(transitions[0].get("advTm"))
        self.assertIsNotNone(transitions[0].find(f"{{{P_NS}}}fade"))

    def test_clear_without_transition_is_noop(self) -> None:
        """没有 transition 时返回 False。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())
        self.assertFalse(_clear_adv_tm(el))

    def test_clear_without_advtm_is_noop(self) -> None:
        """有 transition 但无 advTm 时返回 False，元素保持原样。"""
        xml = (
            f'<p:sld xmlns:p="{P_NS}"><p:cSld/>'
            f'<p:transition spd="slow"/></p:sld>'
        )
        el = lxml_etree.fromstring(xml.encode())
        self.assertFalse(_clear_adv_tm(el))
        trans = el.find(f"{{{P_NS}}}transition")
        self.assertEqual(trans.get("spd"), "slow")

    def test_repeated_clear_is_idempotent(self) -> None:
        """重复清除：第二次返回 False 且无异常。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())
        _set_adv_tm(el, 1000)
        self.assertTrue(_clear_adv_tm(el))
        self.assertFalse(_clear_adv_tm(el))

    def test_set_then_clear_roundtrip(self) -> None:
        """设置后清除恢复未设置状态；再次设置仍正常工作。"""
        xml = f'<p:sld xmlns:p="{P_NS}"><p:cSld/></p:sld>'
        el = lxml_etree.fromstring(xml.encode())

        _set_adv_tm(el, 2000)
        _clear_adv_tm(el)
        self.assertIsNone(el.find(f"{{{P_NS}}}transition"))

        _set_adv_tm(el, 4000)
        trans = el.find(f"{{{P_NS}}}transition")
        self.assertIsNotNone(trans)
        self.assertEqual(trans.get("advTm"), "4000")


class TestAutoAdvance(unittest.IsolatedAsyncioTestCase):
    """自动翻页集成测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _make_config(self, **overrides) -> PptSpeechConfig:
        defaults = dict(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
            auto_advance=True,
            auto_advance_delay=2.0,
        )
        defaults.update(overrides)
        return PptSpeechConfig(**defaults)

    @patch("ppt_speech.core.pipeline.set_advance_after_time")
    @patch("ppt_speech.core.pipeline.get_audio_duration", return_value=5.0)
    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_auto_advance_sets_duration_plus_delay(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
        mock_get_duration: MagicMock,
        mock_set_advance: MagicMock,
    ) -> None:
        """开启 auto_advance 时按「音频时长 + n」设置翻页时间。"""
        mock_read_notes.return_value = "备注"
        mock_text_to_mp3.return_value = True

        config = self._make_config()
        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, config)

        mock_get_duration.assert_called_once()
        # delay = 5.0 + 2.0 = 7.0
        mock_set_advance.assert_called_once_with(mock_slide, 7.0)

    @patch("ppt_speech.core.pipeline.set_advance_after_time")
    @patch("ppt_speech.core.pipeline.get_audio_duration")
    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_auto_advance_disabled_not_called(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
        mock_get_duration: MagicMock,
        mock_set_advance: MagicMock,
    ) -> None:
        """关闭 auto_advance 时不应调用翻页相关函数。"""
        mock_read_notes.return_value = "备注"
        mock_text_to_mp3.return_value = True

        config = self._make_config(auto_advance=False)
        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, config)

        mock_get_duration.assert_not_called()
        mock_set_advance.assert_not_called()

    @patch("ppt_speech.core.pipeline.set_advance_after_time")
    @patch(
        "ppt_speech.pipeline.get_audio_duration",
        side_effect=ValueError("解析失败"),
    )
    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_duration_failure_skips_advance_gracefully(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
        mock_get_duration: MagicMock,
        mock_set_advance: MagicMock,
    ) -> None:
        """音频时长解析失败时优雅跳过，不影响整体流程。"""
        mock_read_notes.return_value = "备注"
        mock_text_to_mp3.return_value = True

        config = self._make_config()
        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, config)  # 不应抛出异常

        mock_set_advance.assert_not_called()
        mock_embed.assert_called_once()  # 嵌入仍正常
        mock_prs.save.assert_called_once()  # 保存仍正常

    @patch("ppt_speech.core.pipeline.set_advance_after_time")
    @patch("ppt_speech.core.pipeline.get_audio_duration")
    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.read_notes_text")
    async def test_custom_delay_used(
        self,
        mock_read_notes: MagicMock,
        mock_text_to_mp3: AsyncMock,
        mock_embed: MagicMock,
        mock_get_duration: MagicMock,
        mock_set_advance: MagicMock,
    ) -> None:
        """自定义 n 值应体现在翻页时间中。"""
        mock_read_notes.return_value = "备注"
        mock_text_to_mp3.return_value = True
        mock_get_duration.return_value = 10.0

        config = self._make_config(auto_advance_delay=0.5)
        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]

        await process_slides(mock_prs, config)

        # delay = 10.0 + 0.5 = 10.5
        mock_set_advance.assert_called_once_with(mock_slide, 10.5)


class TestDubbingIdempotency(unittest.IsolatedAsyncioTestCase):
    """重复配音幂等性集成测试（真实嵌入，仅模拟 TTS/备注/时长）。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.prs = Presentation()
        self.slides = [
            self.prs.slides.add_slide(self.prs.slide_layouts[6])
            for _ in range(2)
        ]

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _make_config(self, **overrides) -> PptSpeechConfig:
        defaults = dict(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
            auto_advance=True,
            auto_advance_delay=2.0,
        )
        defaults.update(overrides)
        return PptSpeechConfig(**defaults)

    @staticmethod
    def _slide_state(slide) -> tuple[int, int, "str | None"]:
        element = slide._element
        pics = len(element.findall(f".//{{{P_NS}}}pic"))
        videos = len(element.findall(f".//{{{P_NS}}}video"))
        transition = element.find(f"{{{P_NS}}}transition")
        adv_tm = (
            transition.get("advTm")
            if transition is not None
            else None
        )
        return pics, videos, adv_tm

    async def _run(self, notes: list[str], config: PptSpeechConfig) -> None:
        async def fake_tts(text, save_path, **_kwargs):
            save_path = Path(save_path)
            save_path.parent.mkdir(parents=True, exist_ok=True)
            save_path.write_bytes(b"ID3" + b"\x00" * 100)
            return True

        with patch(
            "ppt_speech.core.pipeline.read_notes_text",
            side_effect=notes,
        ), patch(
            "ppt_speech.core.pipeline.text_to_mp3",
            new=AsyncMock(side_effect=fake_tts),
        ), patch(
            "ppt_speech.core.pipeline.get_audio_duration",
            return_value=4.0,
        ):
            await process_slides(self.prs, config)

    async def test_repeated_process_is_idempotent(self) -> None:
        """同一演示文稿连续配音两轮：形状/时序不叠加，定时被更新。"""
        config = self._make_config()
        await self._run(["第一页", "第二页"], config)
        with zipfile.ZipFile(config.output_path) as archive:
            mp3_after_first = len(
                [
                    name
                    for name in archive.namelist()
                    if name.startswith("ppt/media/")
                    and name.endswith(".mp3")
                ]
            )

        await self._run(["第一页新", "第二页新"], config)

        for slide in self.slides:
            pics, videos, adv_tm = self._slide_state(slide)
            self.assertEqual(pics, 1)
            self.assertEqual(videos, 1)
            # 4.0s 音频 + 2.0s 缓冲 = 6000ms
            self.assertEqual(adv_tm, "6000")

        with zipfile.ZipFile(config.output_path) as archive:
            mp3_after_second = len(
                [
                    name
                    for name in archive.namelist()
                    if name.startswith("ppt/media/")
                    and name.endswith(".mp3")
                ]
            )
        # 两轮后媒体部件数量不增长（相同字节本就被包级去重）
        self.assertEqual(mp3_after_second, mp3_after_first)

    async def test_notes_removed_on_rerun_strips_old_dubbing(self) -> None:
        """首轮有备注、次轮备注清空：旧配音与旧翻页定时应被一并清除。"""
        config = self._make_config()
        await self._run(["第一页", "第二页"], config)
        for slide in self.slides:
            self.assertEqual(self._slide_state(slide)[0], 1)

        await self._run(["", ""], config)

        for slide in self.slides:
            pics, videos, adv_tm = self._slide_state(slide)
            self.assertEqual(pics, 0)
            self.assertEqual(videos, 0)
            self.assertIsNone(adv_tm)
            self.assertIsNone(
                slide._element.find(f"{{{P_NS}}}timing")
            )

    async def test_disable_autoadvance_on_rerun_clears_old_timing(self) -> None:
        """首轮开启自动翻页、次轮关闭：旧 advTm 被清除，音频保留。"""
        config_on = self._make_config(auto_advance=True)
        await self._run(["第一页", "第二页"], config_on)

        config_off = self._make_config(
            auto_advance=False, output_filename="output2.pptx"
        )
        await self._run(["第一页", "第二页"], config_off)

        for slide in self.slides:
            pics, videos, adv_tm = self._slide_state(slide)
            self.assertEqual(pics, 1)
            self.assertEqual(videos, 1)
            self.assertIsNone(adv_tm)


class TestMain(unittest.IsolatedAsyncioTestCase):
    """main 函数测试。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())
        self.input_path = self.temp_dir / "input.pptx"
        self.input_path.write_bytes(b"fake pptx content")

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    @patch("ppt_speech.core.pipeline.process_slides", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.Presentation")
    async def test_main_with_custom_config(
        self,
        mock_presentation: MagicMock,
        mock_process: AsyncMock,
    ) -> None:
        """测试使用自定义配置运行。"""
        mock_prs = MagicMock()
        mock_presentation.return_value = mock_prs

        config = PptSpeechConfig(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
        )

        await speak_ppt_notes(config)

        mock_presentation.assert_called_once_with(str(self.input_path))
        mock_process.assert_called_once_with(mock_prs, config)

    @patch("ppt_speech.core.pipeline.process_slides", new_callable=AsyncMock)
    @patch("ppt_speech.core.pipeline.Presentation")
    async def test_main_with_default_config(
        self,
        mock_presentation: MagicMock,
        mock_process: AsyncMock,
    ) -> None:
        """测试使用默认配置运行。"""
        mock_prs = MagicMock()
        mock_presentation.return_value = mock_prs

        with patch.object(
            PptSpeechConfig, "validate", return_value=None
        ):
            with patch.object(
                PptSpeechConfig,
                "input_path",
                new_callable=PropertyMock,
                return_value=self.input_path,
            ):
                await speak_ppt_notes()

        mock_presentation.assert_called_once()
        mock_process.assert_called_once()

    async def test_main_propagates_validation_error(self) -> None:
        """测试 main 函数传播验证错误。"""
        config = PptSpeechConfig(
            input_dir=Path("nonexistent_dir"),
        )
        with self.assertRaises(FileNotFoundError):
            await speak_ppt_notes(config)


class TestIntegration(unittest.IsolatedAsyncioTestCase):
    """集成测试：验证模块间协作。"""

    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        if self.temp_dir.exists():
            shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_config_paths_combined_correctly(self) -> None:
        """测试配置路径正确组合。"""
        config = PptSpeechConfig(
            input_dir=Path("inputs"),
            output_dir=Path("outputs"),
            input_filename="presentation.pptx",
            output_filename="result.pptx",
        )
        self.assertEqual(
            config.input_path, Path("inputs/presentation.pptx")
        )
        self.assertEqual(
            config.output_path, Path("outputs/result.pptx")
        )

    def test_voice_normalization_roundtrip(self) -> None:
        """测试语音名称规范化往返。"""
        short_name = "zh-CN-XiaoxiaoNeural"
        full_name = normalize_voice_name(short_name)
        self.assertIn("Microsoft Server Speech Text to Speech Voice", full_name)
        self.assertIn("zh-CN", full_name)
        self.assertIn("XiaoxiaoNeural", full_name)

    def test_voice_normalization_with_region(self) -> None:
        """测试带地区的语音名称规范化。"""
        short_name = "en-US-AriaNeural"
        full_name = normalize_voice_name(short_name)
        self.assertIn("en-US", full_name)
        self.assertIn("AriaNeural", full_name)

    @patch("ppt_speech.core.pipeline.Presentation")
    @patch("ppt_speech.core.pipeline.embed_audio_autoplay")
    @patch("ppt_speech.core.pipeline.read_notes_text")
    @patch("ppt_speech.core.pipeline.text_to_mp3", new_callable=AsyncMock)
    async def test_full_pipeline_mock(
        self,
        mock_text_to_mp3: AsyncMock,
        mock_read_notes: MagicMock,
        mock_embed: MagicMock,
        mock_presentation: MagicMock,
    ) -> None:
        """测试完整流水线（使用 mock）。"""
        mock_read_notes.return_value = "集成测试备注"
        mock_text_to_mp3.return_value = True

        mock_slide = MagicMock()
        mock_prs = MagicMock()
        mock_prs.slides = [mock_slide]
        mock_presentation.return_value = mock_prs

        input_path = self.temp_dir / "input.pptx"
        input_path.write_bytes(b"fake")

        config = PptSpeechConfig(
            input_dir=self.temp_dir,
            output_dir=self.temp_dir,
            input_filename="input.pptx",
            output_filename="output.pptx",
            temp_audio_dir=self.temp_dir / "temp_audio",
            auto_advance=False,
        )

        await speak_ppt_notes(config)

        mock_text_to_mp3.assert_called_once()
        mock_embed.assert_called_once()
        mock_prs.save.assert_called_once_with(str(config.output_path))
        self.assertFalse(config.temp_audio_dir.exists())


if __name__ == "__main__":
    unittest.main()